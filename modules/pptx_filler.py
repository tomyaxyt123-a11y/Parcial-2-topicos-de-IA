"""
Módulo para llenar la plantilla oficial Certificados.pptx y convertir a PDF.
Universidad de La Salle · Tópicos de IA
"""

from __future__ import annotations

import logging
import os
import subprocess
from pathlib import Path
from typing import Dict, List, Optional
from pptx import Presentation

import config
from modules import pdf_generator

LOG = logging.getLogger("certificado")
PLANTILLA_PPTX_OFICIAL = config.BASE_DIR / "Certificados.pptx"


def llenar_pptx(
    ruta_pptx: str,
    datos: Dict[str, object],
    ruta_salida_pptx: str,
) -> str:
    """Abre la plantilla PPTX, reemplaza los placeholders institucionales y guarda el archivo."""
    prs = Presentation(ruta_pptx)
    slide = prs.slides[0]

    nombre = str(datos.get("nombre_completo") or datos.get("nombres") or "Participante").strip()
    rol = str(datos.get("rol") or "Participante").strip()
    aporte = str(datos.get("ponencia") or datos.get("conferencia") or "").strip()
    modalidad = f"{rol} · «{aporte}»" if aporte else rol

    reemplazos = {
        "<NOMBRE>": nombre,
        "<PROMEDIO>": "CERTIFICADO",
        "<CURSO>": modalidad,
        "OTORTGA": "OTORGA",
        "Faculdad": "Facultad",
    }

    for shape in slide.shapes:
        if shape.has_text_frame:
            for p in shape.text_frame.paragraphs:
                for run in p.runs:
                    for k, v in reemplazos.items():
                        if k in run.text:
                            run.text = run.text.replace(k, v)

    Path(ruta_salida_pptx).parent.mkdir(parents=True, exist_ok=True)
    prs.save(ruta_salida_pptx)
    return str(ruta_salida_pptx)


def convertir_lote_pptx_a_pdf(rutas_pptx: List[str]) -> List[str]:
    """Convierte una lista de archivos PPTX a PDF usando PowerPoint COM en un solo proceso headless."""
    rutas_pendientes = [
        str(Path(p).resolve()) for p in rutas_pptx
        if Path(p).exists() and not Path(p).with_suffix(".pdf").exists()
    ]
    if not rutas_pendientes:
        return [str(Path(p).with_suffix(".pdf")) for p in rutas_pptx]

    # Convertir por PowerShell COM en lote para máxima velocidad
    script_ps1 = config.BASE_DIR / "scripts" / "convertir_lote.ps1"
    lista_txt = config.BASE_DIR / "scripts" / "lista_rutas_conversion.txt"
    script_ps1.parent.mkdir(parents=True, exist_ok=True)
    lista_txt.write_text("\n".join(rutas_pendientes), encoding="utf-8")

    contenido_script = """\
param([string]$listaPath)
if (-not (Test-Path $listaPath)) { exit 0 }
$rutas = Get-Content $listaPath -Encoding UTF8 | Where-Object { $_.Trim() -ne "" }
if ($rutas.Count -eq 0) { exit 0 }

$ppt = New-Object -ComObject PowerPoint.Application
try {
    foreach ($archivo in $rutas) {
        $archivo = $archivo.Trim()
        if (Test-Path $archivo) {
            $pdfPath = [System.IO.Path]::ChangeExtension($archivo, "pdf")
            if (-not (Test-Path $pdfPath)) {
                $pres = $ppt.Presentations.Open((Resolve-Path $archivo).Path, 1, 0, 0)
                $pres.SaveAs($pdfPath, 32)
                $pres.Close()
            }
        }
    }
} finally {
    $ppt.Quit()
    [System.Runtime.Interopservices.Marshal]::ReleaseComObject($ppt) | Out-Null
    [System.GC]::Collect()
    [System.GC]::WaitForPendingFinalizers()
}
"""
    script_ps1.write_text(contenido_script, encoding="utf-8")

    try:
        cmd = [
            "powershell",
            "-NoProfile",
            "-NonInteractive",
            "-ExecutionPolicy", "Bypass",
            "-File", str(script_ps1),
            "-listaPath", str(lista_txt),
        ]
        res = subprocess.run(cmd, capture_output=True, text=True, timeout=180)
        if res.returncode != 0:
            LOG.warning("[PPTX] Error en conversión COM: %s", res.stderr)
    except Exception as exc:
        LOG.warning("[PPTX] Excepción ejecutando conversión COM: %s", exc)
    finally:
        try:
            lista_txt.unlink(missing_ok=True)
        except OSError:
            pass

    return [str(Path(p).with_suffix(".pdf")) for p in rutas_pptx]


def generar_pptx_incietec(
    datos: Dict[str, object],
    carpeta_salida: Path,
    slug: str,
) -> Dict[str, str]:
    """Genera el archivo PPTX llenando los campos sin convertir aún a PDF (ultra rápido)."""
    nombre_base = f"{slug}"
    ruta_pptx = carpeta_salida / f"{nombre_base}.pptx"
    ruta_pdf = carpeta_salida / f"{nombre_base}.pdf"

    llenar_pptx(str(PLANTILLA_PPTX_OFICIAL), datos, str(ruta_pptx))

    return {
        "pptx": str(ruta_pptx),
        "pdf": str(ruta_pdf),
        "nombre_pdf": ruta_pdf.name,
        "nombre_pptx": ruta_pptx.name,
    }


def generar_certificado_incietec(
    datos: Dict[str, object],
    carpeta_salida: Path,
    slug: str,
    proteger: bool = False,
    contrasena: Optional[str] = None,
) -> Dict[str, str]:
    """Genera el certificado a partir de Certificados.pptx llenando los campos."""
    res = generar_pptx_incietec(datos, carpeta_salida, slug)
    ruta_pptx = Path(res["pptx"])
    ruta_pdf = Path(res["pdf"])

    # Convertir a PDF si no existe
    convertir_lote_pptx_a_pdf([str(ruta_pptx)])

    # Si por alguna razón la conversión falló, fallback con plantilla base
    if not ruta_pdf.exists():
        plantilla_fondo = config.TEMPLATES_DIR / "Certificados.pdf"
        if not plantilla_fondo.exists():
            plantilla_fondo = config.TEMPLATES_DIR / "certificado_base.pdf"
        pdf_generator.generar_certificado(
            datos,
            str(plantilla_fondo) if plantilla_fondo.exists() else None,
            str(ruta_pdf),
            proteger=proteger,
            contrasena=contrasena,
        )
    elif proteger:
        clave = contrasena or str(datos.get("documento") or "").strip()
        if clave:
            pdf_generator._cifrar_pdf(str(ruta_pdf), clave)

    return res
